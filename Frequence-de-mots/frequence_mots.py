# by Renso93
# www.github.com/renso93
#   
import re
from collections import Counter
from unidecode import unidecode
from PyPDF2 import PdfReader
from docx import Document
from pathlib import Path


class WordFrequencyAnalyzer:
    stop_words_raw = {
        "l", "le", "la", "les", "a", "des", "un", "une", "de", "d",
        "du", "au", "en", "et", "est", "à", "par", "pour", "sur", "avec",
    }

    def __init__(self, file_path:str):
        self.file_path = file_path
        self.stop_words = {unidecode(word).lower() for word in self.stop_words_raw}

    def read_pdf(self):
        """Lit un fichier PDF et retourne son contenu sous forme de générateur."""
        try:
            with open(self.file_path, "rb") as file:
                pdf = PdfReader(file)
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        yield page_text
        except Exception as e:
            print(f"Erreur lors de la lecture du PDF : {e}")
            return

    def read_docx(self):
        """Lit un fichier DOCX et retourne son contenu sous forme de générateur."""
        try:
            doc = Document(self.file_path)
            for paragraph in doc.paragraphs:
                if paragraph.text and paragraph.text.strip():
                    yield paragraph.text

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text and cell.text.strip():
                            yield cell.text
        except Exception as e:
            print(f"Erreur lors de la lecture du DOCX : {e}")
            return

    def read_file(self):
        """Lit un fichier (PDF, DOCX ou texte) et retourne son contenu."""
        suffix = Path(self.file_path).suffix.lower()
        if suffix.endswith(".pdf"):
            yield from self.read_pdf()
        if suffix.endswith(".docx"):
            yield from self.read_docx()

        try:
            with open(self.file_path, "r", encoding="utf-8") as f:
                for line in f:
                    if line.strip():
                        yield line
        except FileNotFoundError:
            print(f"Fichier introuvable : {self.file_path}")
            return
        except Exception as e:
            print(f"Erreur lors de la lecture du fichier : {e}")
            return

    def clean_text(self, text: str) -> str:
        """Nettoie le texte en supprimant la ponctuation et en convertissant en minuscules."""
        text = unidecode(text)
        return re.sub(r"[^\w\s]", " ", text.lower())

    def filter_words(self, words):
        """Filtre les mots en supprimant les mots vides et les mots trop courts."""
        return [word for word in words if word not in self.stop_words and len(word) > 1]

    def analyze_frequency(self, n=5):
        """Analyse la fréquence des mots dans le fichier."""
        content = self.read_file()
        if content is None:
            return []

        content_list = list(content)
        if not content_list or all(item is None for item in content_list):
            print("Aucun contenu valide extrait du fichier.")
            return []

        text = " ".join(content_list)
        cleaned_text = self.clean_text(text)
        words = cleaned_text.split()
        filtered_words = self.filter_words(words)
        frequency = Counter(filtered_words)
        return frequency.most_common(n)


if __name__ == "__main__":
    file_path = "file/Pierre-Giraud-Liste-des-propriétés-CSS.pdf"  # Remplacez par le chemin de votre fichier
    analyzer = WordFrequencyAnalyzer(file_path=file_path)
    top_words = analyzer.analyze_frequency(10)

    print(f"\n== LES {len(top_words)} MOTS LES PLUS FRÉQUENTS ==\n")
    for word, count in top_words:
        print(f"{word}: {count}")